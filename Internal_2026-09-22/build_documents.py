from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
ROLL = "2311CS040156"
NAME = "Ataur"
DATE = "2026-09-22"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 2"].font.bold = True


def add_meta_table(document):
    table = document.add_table(rows=3, cols=2)
    table.autofit = True
    data = [("Name", NAME), ("Roll number", ROLL), ("Date", DATE)]
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if cell_index == 0:
                cell.text = data[row_index][0]
                set_cell_shading(cell, "D9EAF7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            else:
                cell.text = data[row_index][1]
    document.add_paragraph()


def add_paragraph(document, text, bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        run.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def create_a4_report():
    document = Document()
    style_document(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Ethical Hacking Internal Report")
    add_meta_table(document)

    document.add_heading("Question 5 SQL Injection Detection and Prevention", level=1)
    document.add_heading("Aim", level=2)
    add_paragraph(document, "To understand how SQL injection happens when untrusted input is directly joined into a database query, and to show how input validation and parameterized queries prevent it.")
    document.add_heading("Key Concepts", level=2)
    add_paragraph(document, "SQL injection is a web application weakness where user input is interpreted as part of an SQL command. It can allow an attacker to bypass authentication, read unintended records, modify data, or delete data. Detection usually involves checking whether special SQL characters or boolean payloads change the behavior of a form. Prevention depends on parameterized queries, input validation, least privilege database accounts, and safe error handling.")
    document.add_heading("Lab Target", level=2)
    add_paragraph(document, "Suggested target: DVWA. Submission method used here: safe local simulation using the same idea as a vulnerable login or search form. No external target was tested.")
    document.add_heading("Process", level=2)
    add_bullets(document, [
        "A sample user table was considered with normal records such as admin and student.",
        "A vulnerable query pattern was reviewed where the application builds a query by directly adding user input into the SQL string.",
        "A test input such as ' OR '1'='1' -- was used in the simulation to represent the injection attempt.",
        "The same input was then checked with a parameterized query, where the input is treated as data only.",
        "Basic input validation and error handling practices were noted as additional controls.",
    ])
    document.add_heading("Result", level=2)
    add_paragraph(document, "The vulnerable query treated the injected condition as SQL logic and returned a successful match in the simulated test. The parameterized query treated the same payload as ordinary text and returned no unauthorized match. This shows that parameterized queries are the main defense against SQL injection, while validation and least privilege reduce additional risk.")
    document.add_heading("Prevention Points", level=2)
    add_bullets(document, [
        "Use parameterized queries or prepared statements for every database operation.",
        "Validate input length, type, and expected format before processing it.",
        "Avoid displaying raw SQL errors to users.",
        "Use a database account with only the permissions required by the application.",
        "Log suspicious inputs for review.",
    ])

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Question 7 Password Security and Authentication Testing", level=1)
    document.add_heading("Aim", level=2)
    add_paragraph(document, "To study password policy weaknesses, secure password storage, salting, hashing, and common authentication risks.")
    document.add_heading("Key Concepts", level=2)
    add_paragraph(document, "Password security depends on both user password strength and server-side storage. Weak passwords can be guessed easily, and plain text password storage can expose all accounts if the database is leaked. Secure systems store salted password hashes using slow hashing functions. Authentication testing checks password policy, login error messages, account lockout, session handling, and resistance to brute-force attempts.")
    document.add_heading("Lab Target", level=2)
    add_paragraph(document, "Suggested target: local test application. Submission method used here: safe local simulation of password validation, salted hashing, and login lockout logic.")
    document.add_heading("Process", level=2)
    add_bullets(document, [
        "Several sample passwords were checked against a basic password policy.",
        "The policy checked minimum length, uppercase letters, lowercase letters, digits, and special characters.",
        "A strong password was converted into a salted hash using a slow key derivation function.",
        "A login verification step compared the submitted password hash with the stored hash.",
        "Repeated failed login attempts were simulated to demonstrate account lockout behavior.",
    ])
    document.add_heading("Result", level=2)
    add_paragraph(document, "Weak passwords failed the policy checks because they were short or lacked character variety. The strong password passed the policy check and was stored as a salted hash instead of plain text. During authentication, the correct password verified successfully, while incorrect attempts failed. After repeated failures, the account was locked in the simulation. This shows that secure authentication requires strong password rules, salted slow hashing, generic error messages, and rate limiting or lockout controls.")
    document.add_heading("Prevention Points", level=2)
    add_bullets(document, [
        "Never store passwords in plain text.",
        "Use a unique salt for every password.",
        "Use slow password hashing such as PBKDF2, bcrypt, scrypt, or Argon2.",
        "Apply rate limiting or temporary lockout after repeated failed attempts.",
        "Use generic login errors such as Invalid username or password.",
        "Enable multi-factor authentication where possible.",
    ])
    document.add_heading("Conclusion", level=1)
    add_paragraph(document, "SQL injection and weak authentication are common application security risks. SQL injection is prevented mainly by parameterized queries and careful input handling. Password attacks are reduced by strong password policies, salted slow hashing, and login protection controls. The simulations show the difference between unsafe and safe implementation patterns without testing any real external system.")
    path = ROOT / "{}_A4_Report.docx".format(ROLL)
    document.save(path)
    return path


def add_code_block(document, code):
    for line in code.strip("\n").splitlines():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)


def create_code_result_doc():
    document = Document()
    style_document(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Code and Result Document")
    add_meta_table(document)
    add_paragraph(document, "This document contains safe local demonstration code and results for the assigned internal questions. The output represents controlled local simulations, not testing of any live or external target.")

    document.add_heading("Question 5 SQL Injection Detection and Prevention", level=1)
    document.add_heading("Code", level=2)
    add_code_block(document, (ROOT / "sql_injection_demo.py").read_text(encoding="utf-8"))
    document.add_heading("Result", level=2)
    add_code_block(document, """SQL Injection Detection and Prevention Demo
Payload username: ' OR '1'='1' --
Vulnerable query: SELECT username FROM users WHERE username = '' OR '1'='1' -- ' AND password = 'anything'
Vulnerable result: [('admin',), ('student',)]
Safe query: SELECT username FROM users WHERE username = ? AND password = ?
Safe result: []""")
    document.add_heading("Interpretation", level=2)
    add_paragraph(document, "The vulnerable query joined user input directly into SQL, so the payload changed the query logic. The safe query used placeholders, so the database treated the payload as text and did not return unauthorized rows.")

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Question 7 Password Security and Authentication Testing", level=1)
    document.add_heading("Code", level=2)
    add_code_block(document, (ROOT / "password_security_demo.py").read_text(encoding="utf-8"))
    document.add_heading("Result", level=2)
    add_code_block(document, """Password Security and Authentication Testing Demo
admin {'minimum_length': False, 'uppercase': False, 'lowercase': True, 'digit': False, 'special_character': False} FAIL
password123 {'minimum_length': True, 'uppercase': False, 'lowercase': True, 'digit': True, 'special_character': False} FAIL
Student@2026 {'minimum_length': True, 'uppercase': True, 'lowercase': True, 'digit': True, 'special_character': True} PASS
Salt length: 16
Hash length: 32
Correct password verification: True
Wrong password verification: False
Lockout result: Account locked after repeated failed attempts""")
    document.add_heading("Interpretation", level=2)
    add_paragraph(document, "The weak passwords failed because they did not satisfy the full policy. The strong password passed and was stored as a salted PBKDF2 hash. Verification succeeded only for the correct password, and the lockout simulation blocked repeated failed attempts.")

    path = ROOT / "{}_Code_and_Result.docx".format(ROLL)
    document.save(path)
    return path


if __name__ == "__main__":
    print(create_a4_report())
    print(create_code_result_doc())
